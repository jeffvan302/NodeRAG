#!/usr/bin/env python3
"""Generate the M12 DOCX and PDF projections of the real NodeRag planning corpus."""

from __future__ import annotations

import argparse
from io import BytesIO
from pathlib import Path


def output_directory() -> Path:
    path = Path(__file__).resolve().parents[1] / "fixtures" / "evaluation" / "m12" / "corpus"
    path.mkdir(parents=True, exist_ok=True)
    return path


def create_docx() -> Path:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    target = output_directory() / "project-requirements.docx"
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    document.core_properties.title = "NodeRag Project Requirements"
    document.core_properties.author = "NodeRag Project"
    document.core_properties.subject = "Milestone 12 evaluation corpus"

    for style_name, size, bold in [
        ("Normal", 11, False),
        ("Title", 25, True),
        ("Heading 1", 17, True),
        ("Heading 2", 12, True),
    ]:
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.paragraph_format.space_after = Pt(7 if style_name == "Normal" else 9)
        style.paragraph_format.line_spacing = 1.08
        if style_name == "Title":
            style_properties = style._element.get_or_add_pPr()
            style_border = style_properties.find(qn("w:pBdr"))
            if style_border is not None:
                style_properties.remove(style_border)

    def shade(cell, color: str) -> None:
        properties = cell._tc.get_or_add_tcPr()
        fill = properties.find(qn("w:shd"))
        if fill is None:
            fill = OxmlElement("w:shd")
            properties.append(fill)
        fill.set(qn("w:fill"), color)

    def border_table(table) -> None:
        properties = table._tbl.tblPr
        borders = properties.find(qn("w:tblBorders"))
        if borders is None:
            borders = OxmlElement("w:tblBorders")
            properties.append(borders)
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            item = OxmlElement(f"w:{edge}")
            item.set(qn("w:val"), "single")
            item.set(qn("w:sz"), "4")
            item.set(qn("w:color"), "D9D9D9")
            borders.append(item)

    def style_table(table, widths: list[float]) -> None:
        table.autofit = False
        border_table(table)
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                cell.width = Inches(widths[column_index])
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                cell.margin_top = Inches(0.08)
                cell.margin_bottom = Inches(0.08)
                cell.margin_left = Inches(0.09)
                cell.margin_right = Inches(0.09)
                if row_index == 0:
                    shade(cell, "1F4E78")
                elif row_index % 2 == 0:
                    shade(cell, "EEF3F8")
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    for run in paragraph.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(9.5)
                        if row_index == 0:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)

    def add_bullets(items: list[str]) -> None:
        for item in items:
            paragraph = document.add_paragraph(item, style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(4)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("NodeRag Project Requirements")
    title_properties = title._p.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    subtitle = document.add_paragraph()
    subtitle.add_run("Approved project decisions for evaluation and retrieval testing").italic = True

    document.add_heading("Deployment and ingestion", level=1)
    document.add_paragraph(
        "NodeRag is a self-contained retrieval augmented generation database for Synology DSM 7.2 or later. "
        "Development and milestone certification run in WSL and Docker through Milestone 14. The exact release "
        "candidate is tested on DSM only after the Docker milestones are complete."
    )
    document.add_paragraph(
        "The package contains PostgreSQL 17 with pgvector and uses Qwen3-Embedding-0.6B on the device CPU. "
        "Document text, embeddings, queries, and search results remain local. Only selected images may be sent "
        "to the configured Ollama vision host."
    )
    document.add_heading("Supported source formats", level=2)
    add_bullets([
        "PDF, DOCX, XLSX, Markdown, and plain text are supported release formats.",
        "Every original and normalized artifact is identified by a SHA-256 digest.",
        "ING-417 identifies the rule that a failed image analysis must not stop text ingestion.",
        "Afrikaans source phrase: plaaslike kennis bly op die NAS.",
        "Japanese source phrase: ローカル検索は文書の境界を守ります。",
    ])

    document.add_page_break()
    document.add_heading("Retrieval quality and locators", level=1)
    document.add_paragraph(
        "Vector search is the initial release default. Hybrid becomes the default only when it improves overall "
        "Recall at 10 by at least two percentage points without more than doubling p95 latency. Keyword search "
        "remains available for exact identifiers, names, error codes, and quoted phrases."
    )
    thresholds = document.add_table(rows=1, cols=3)
    thresholds.rows[0].cells[0].text = "Measure"
    thresholds.rows[0].cells[1].text = "Release threshold"
    thresholds.rows[0].cells[2].text = "Scope"
    for values in [
        ("Overall Recall at 10", "At least 0.85", "Answerable golden cases"),
        ("Format Recall at 10", "At least 0.80", "Each supported format"),
        ("Zero result rate", "No more than 5 percent", "Answerable cases"),
        ("Idle search p95", "No more than 3 seconds", "Recorded Docker envelope"),
        ("Active ingest p95", "No more than 5 seconds", "Recorded Docker envelope"),
    ]:
        cells = thresholds.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    style_table(thresholds, [2.0, 1.8, 3.0])
    document.add_paragraph(
        "A relevant judgment uses the stable document identity plus a page number or sheet range. Chunk IDs are "
        "not relevance identities because re-chunking replaces them. Chunks may span pages and therefore carry "
        "both page_start and page_end as well as UTF-16 character offsets and UTF-8 byte offsets."
    )

    document.add_page_break()
    document.add_heading("Rebuild activation and rollback", level=1)
    document.add_paragraph(
        "A build freezes the chunk profile, model manifest digest, vector dimension, and selected silos. An "
        "administrator confirms a capacity-valid proposal before detached staging tables and low-priority jobs "
        "are created. Search never reads a staging table before verification succeeds."
    )
    lifecycle = document.add_table(rows=1, cols=3)
    lifecycle.rows[0].cells[0].text = "Stage"
    lifecycle.rows[0].cells[1].text = "Required state"
    lifecycle.rows[0].cells[2].text = "Failure behavior"
    for values in [
        ("Build", "Detached and checkpointed", "Resume from completed documents"),
        ("Verify", "Coverage and locator checks pass", "Leave staging unsearchable"),
        ("Activate", "Exclusive cutover lock", "Keep prior active tables"),
        ("Rollback", "Prior table is retained", "Restore before cleanup"),
        ("Cleanup", "Typed authorization exists", "Retry transactionally"),
    ]:
        cells = lifecycle.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = value
    style_table(lifecycle, [1.25, 2.35, 3.2])
    document.add_paragraph(
        "REB-409 means another activation is rejected while an ambiguous rollback chain exists. Search holds a "
        "shared advisory lock across model resolution, query embedding, and retrieval. Activation takes the "
        "matching exclusive lock, which prevents an old-model query vector from reaching new-model rows."
    )

    document.add_page_break()
    document.add_heading("Security and DSM certification", level=1)
    document.add_paragraph(
        "Silos are access-control boundaries. Each request resolves a principal and authorized silo set, and "
        "PostgreSQL row-level security is a backstop. Cross-silo leakage must remain exactly zero. Search-only API "
        "keys cannot download originals or fetch complete Markdown."
    )
    document.add_paragraph(
        "The bundled PostgreSQL server listens on a package-private Unix socket by default. The socket directory "
        "uses mode 0700, database credentials are package controlled, and private filesystem paths are never "
        "returned by the public API. SEC-403 identifies a denied cross-silo retrieval attempt."
    )
    document.add_heading("Final device checks", level=2)
    add_bullets([
        "Load the exact release candidate with Synology Node.js 22 and the DSM C++ runtime.",
        "Verify Package Center install, upgrade, restart, reboot, uninstall, and retained reinstall behavior.",
        "Measure NAS memory, temperature, DSM administration, file sharing, and storage I/O during background work.",
        "Confirm ACL behavior on real ext4 or Btrfs volumes and encrypted shared folders.",
    ])
    document.add_paragraph(
        "Docker evidence is never represented as installed DSM evidence. The final DSM test closes only the "
        "hardware and Package Center checks that cannot be executed inside WSL or a Linux container."
    )

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("NodeRag M12 evaluation corpus")
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(89, 89, 89)
    document.save(target)
    return target


def create_pdf() -> Path:
    from PIL import Image as PillowImage
    from PIL import ImageDraw, ImageFont
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        Image,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    target = output_directory() / "architecture-evidence.pdf"
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleBlack",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=23,
        leading=28,
        textColor=colors.black,
        alignment=TA_LEFT,
        spaceAfter=12,
    )
    heading_style = ParagraphStyle(
        "HeadingBlack",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor=colors.black,
        spaceAfter=10,
    )
    body_style = ParagraphStyle(
        "BodyReadable",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10.5,
        leading=15,
        textColor=colors.black,
        spaceAfter=8,
    )
    small_style = ParagraphStyle(
        "SmallReadable",
        parent=body_style,
        fontSize=9,
        leading=12,
    )

    def footer(canvas, document) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#595959"))
        canvas.drawCentredString(letter[0] / 2, 0.38 * inch, f"NodeRag architecture evidence  |  {document.page}")
        canvas.restoreState()

    def table(rows, widths):
        result = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
        result.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E78")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                    ("LEADING", (0, 0), (-1, -1), 11),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D9D9D9")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EEF3F8")]),
                    ("LEFTPADDING", (0, 0), (-1, -1), 7),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        return result

    story = [
        Paragraph("NodeRag Architecture Evidence", title_style),
        Paragraph("Approved architecture facts used by the Milestone 12 retrieval evaluation", body_style),
        Spacer(1, 8),
        Paragraph("Local service boundary", heading_style),
        Paragraph(
            "NodeRag runs as one package with a dedicated package user. PostgreSQL 17 and pgvector are bundled "
            "inside the SPK. Qwen3-Embedding-0.6B runs on the NAS CPU and emits 1024-dimensional vectors. The "
            "browser interface, REST API, MCP endpoint, worker, embedder, and database share one authorization "
            "model rather than private UI-only operations.",
            body_style,
        ),
        table(
            [
                ["Component", "Boundary", "Release decision"],
                ["PostgreSQL", "Package-private Unix socket", "TCP is disabled by default"],
                ["Embedder", "Local framed socket", "No runtime model execution off device"],
                ["Vision", "Configured Ollama host", "Only selected images may leave the NAS"],
                ["Gateways", "Ports 5300 through 5319", "Each listener has a fixed silo set"],
            ],
            [1.45 * inch, 2.05 * inch, 2.55 * inch],
        ),
        Spacer(1, 10),
        Paragraph(
            "ARCH-5300 is the exact identifier for the first narrowed gateway port. The control interface uses "
            "port 8787 when DSM has not assigned another package port.",
            body_style,
        ),
        PageBreak(),
        Paragraph("Storage and document flow", heading_style),
        Paragraph(
            "The DSM-selected package volume is the safe default. Bulk chunks, vectors, indexes, models, and "
            "extraction caches may be relocated only after a write, fsync, rename, read, capacity, filesystem, and "
            "permission probe succeeds. PGDATA is not placed in a shared folder used by synchronization or media "
            "indexing services.",
            body_style,
        ),
        Paragraph(
            "Documents are normalized to NFC before offsets and hashes are computed. PDF and Word pages retain "
            "page numbers. Each Excel sheet is treated as an addressable page. Markdown and text use document "
            "character offsets when no page exists. A chunk may cross a page boundary, so page_start and page_end "
            "are both stored.",
            body_style,
        ),
        table(
            [
                ["Artifact", "Integrity rule", "Recovery behavior"],
                ["Original file", "SHA-256 must match", "Report missing or changed source"],
                ["Canonical Markdown", "Write, fsync, atomic rename", "Re-extract from verified original"],
                ["Embedding cache", "Text and model digest key", "Recompute only exact cache misses"],
                ["Vision cache", "Image hash and model key", "Retry without blocking text ingestion"],
            ],
            [1.55 * inch, 2.15 * inch, 2.35 * inch],
        ),
        Spacer(1, 10),
        Paragraph(
            "STOR-0700 identifies the required package-private database socket and data-directory mode. A failed "
            "relocation restores the previous path and records an administrator-visible structured error.",
            body_style,
        ),
        PageBreak(),
        Paragraph("Vision recovery evidence", heading_style),
        Paragraph(
            "This page contains a raster-only project requirement below. The text inside the image is not part of "
            "the PDF text layer. A successful vision path inserts its Markdown at this page locator and records the "
            "provider, model, image digest, confidence, and elapsed time.",
            body_style,
        ),
    ]

    image = PillowImage.new("RGB", (1500, 620), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.load_default(size=34)
    small_font = ImageFont.load_default(size=27)
    draw.rounded_rectangle((35, 35, 1465, 585), radius=26, outline="#1F4E78", width=8, fill="#EEF3F8")
    draw.rectangle((85, 95, 360, 245), outline="#1F4E78", width=6, fill="white")
    draw.text((122, 140), "PDF IMAGE", fill="black", font=font)
    draw.line((390, 170, 635, 170), fill="#1F4E78", width=10)
    draw.polygon([(635, 150), (690, 170), (635, 190)], fill="#1F4E78")
    draw.rectangle((700, 95, 1410, 245), outline="#1F4E78", width=6, fill="white")
    draw.text((755, 125), "VISION MARKDOWN", fill="black", font=font)
    draw.text((755, 180), "page locator and confidence", fill="black", font=small_font)
    draw.text((90, 315), "VR-204", fill="#1F4E78", font=font)
    draw.text((90, 380), "Vision fallback must preserve the page locator", fill="black", font=font)
    draw.text((90, 435), "and confidence score.", fill="black", font=font)
    draw.text((90, 515), "A provider failure leaves the image pending and text searchable.", fill="black", font=small_font)
    image_bytes = BytesIO()
    image.save(image_bytes, format="PNG")
    image_bytes.seek(0)
    raster = Image(image_bytes, width=6.5 * inch, height=2.69 * inch)
    story.extend(
        [
            raster,
            Spacer(1, 12),
            Paragraph(
                "Decorative images below the configured size threshold are skipped. Repeated images are deduplicated "
                "globally by content hash, while cache reference counts ensure that a destructive purge removes "
                "vision text that no remaining document references.",
                body_style,
            ),
            PageBreak(),
            Paragraph("Evaluation and performance rules", heading_style),
            Paragraph(
                "The golden set contains real questions tied to stable document and location identities. Each run "
                "pins the commit, migration head, active build, model manifest, settings revision, corpus snapshot, "
                "and search mode. Reports include Recall at 10, mean reciprocal rank, nDCG at 10, zero-result rate, "
                "and p50 and p95 latency overall and by the required dimensions.",
                body_style,
            ),
            table(
                [
                    ["Decision", "Acceptance rule"],
                    ["Cross-silo leakage", "Exactly zero"],
                    ["Locator mismatch", "Exactly zero"],
                    ["Q4 promotion", "At least 1.5 times faster and no more than 0.01 Recall loss"],
                    ["Hybrid promotion", "At least 0.02 Recall gain and no more than 2 times p95"],
                    ["Chunk RLS", "No more than 15 percent p95 overhead with correct plans"],
                ],
                [2.0 * inch, 4.05 * inch],
            ),
            Spacer(1, 10),
            Paragraph(
                "PERF-3000 is the provisional idle search p95 limit in milliseconds. PERF-5000 is the equivalent "
                "limit while background ingest is active. These values qualify only the recorded Docker resource "
                "envelope and do not become NAS performance promises until final DSM certification.",
                body_style,
            ),
            Paragraph(
                "The nightly run compares the accepted baseline with the current candidate. It alerts on quality, "
                "latency, locator, partition coverage, model-version, or cross-silo regressions.",
                small_style,
            ),
        ]
    )

    document = SimpleDocTemplate(
        str(target),
        pagesize=letter,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.65 * inch,
        bottomMargin=0.65 * inch,
        title="NodeRag Architecture Evidence",
        author="NodeRag Project",
        subject="Milestone 12 evaluation corpus",
    )
    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return target


def main() -> None:
    parser = argparse.ArgumentParser()
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--docx", action="store_true")
    group.add_argument("--pdf", action="store_true")
    options = parser.parse_args()
    target = create_docx() if options.docx else create_pdf()
    print(target)


if __name__ == "__main__":
    main()
